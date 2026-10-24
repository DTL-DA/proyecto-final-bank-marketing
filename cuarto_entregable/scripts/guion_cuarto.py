"""Reescribe los criterios 1 a 7 del guion de sustentación para el cuarto entregable (modelo original L2E).

Conserva el encabezado, la tabla de tiempos y los recuadros «Tiempo y ponderación» de cada criterio de la rúbrica;
reemplaza el guion de cada uno y sus cuadros con las cifras que exporta el cuaderno en ``resultados/guion``. Después
regenera la sección 8 (mejoras) con ``actualizar_guion.py`` para que la numeración de los cuadros sea consecutiva.

Uso (desde la raíz del proyecto, después de ejecutar el cuaderno):
    python scripts/guion_cuarto.py
"""

import json
import pickle
import re
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "scripts"))

import docx  # noqa: E402
import pandas as pd  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from guion_docx import Guion  # noqa: E402

GUION = RAIZ / "Guion de sustentación - Rúbrica - Bank Marketing.docx"
RES = RAIZ / "resultados"
CUAD = "«Proyecto Final - Cuarto Entregable - Bank Marketing.ipynb»"
NOMBRES = {"naive_bayes": "Naive Bayes", "logistica": "la logística L1/L2", "bosque": "Random Forest", "svm": "la SVM",
           "xgboost": "XGBoost", "arbol": "el árbol de decisión", "knn": "KNN"}


def num(x, d=3):
    return f"{x:,.{d}f}".replace(",", " ").replace(".", ",")


def pv(p):
    return "< 0,001" if p < 0.001 else num(p, 3)


def cargar():
    g = RES / "guion"
    D = {n: json.loads((g / f"{n}.json").read_text(encoding="utf-8"))
         for n in ("eda", "literatura", "pruebas_unitarias", "complejidad", "shap", "robustez", "correcciones",
                   "ablacion_sp")}
    D.update({n: pd.read_parquet(g / f"{n}.parquet")
              for n in ("solo_nivel", "hiperparametros", "ablacion", "resumen", "conteo_r3", "pruebas_r3",
                        "adicionales_rangos")})
    D["r3"] = json.loads((RES / "mejoras" / "resumen_r3.json").read_text(encoding="utf-8"))
    D["flex"] = pd.read_parquet(RES / "mejoras" / "flexibles.parquet")
    D["rango"] = pd.read_parquet(RES / "mejoras" / "rango_macro.parquet").set_index("Variable")
    with open(RES / "finales_original" / "l2e.pkl", "rb") as f:
        D["params"] = {k.split("__")[1]: v for k, v in pickle.load(f)["params"].items()}
    return D


def _texto(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _es_h1(doc, el):
    if el.tag != qn("w:p"):
        return False
    ps = el.find(f".//{qn('w:pStyle')}")
    return ps is not None and ps.get(qn("w:val")) == doc.styles["Heading 1"].style_id


def _poner_texto(parrafo_el, texto):
    """Reemplaza el texto de un párrafo conservando el formato de su primer run."""
    runs = parrafo_el.findall(qn("w:r"))
    for r in runs[1:]:
        parrafo_el.remove(r)
    t = runs[0].find(qn("w:t"))
    t.text = texto
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _celda(tabla_el, fila, col, texto):
    tr = tabla_el.findall(qn("w:tr"))[fila]
    tc = tr.findall(qn("w:tc"))[col]
    _poner_texto(tc.find(qn("w:p")), texto)


def limpiar_criterio(doc, k):
    """Deja el título y el recuadro del criterio k y devuelve el recuadro (ancla para el contenido nuevo)."""
    body = doc.element.body
    hijos = list(body.iterchildren())
    i = next(j for j, el in enumerate(hijos) if _es_h1(doc, el) and _texto(el).startswith(f"{k}. "))
    recuadro = hijos[i + 1]
    assert recuadro.tag == qn("w:tbl")
    for el in hijos[i + 2:]:
        if el.tag == qn("w:sectPr") or _es_h1(doc, el):
            break
        body.remove(el)
    return recuadro


def criterio_1(g, D):
    e, r3 = D["eda"], D["r3"]
    g.titulo("Guion", 2)
    g.parrafo("Un banco portugués vende depósitos a término por teléfono y necesita decidir, **antes de marcar**, a quién "
              "llamar primero. El reto es que la tasa de suscripción no se queda quieta: en el entrenamiento "
              f"({e['desde_tr']} a {e['hasta_tr']}) suscribe el {num(100 * e['prev_tr'], 1)} % de los clientes y en la "
              f"prueba ({e['desde_te']} a {e['hasta_te']}), el {num(100 * e['prev_te'], 1)} %. En el tercer entregable "
              "comparamos 140 combinaciones de siete clasificadores estáticos, que asumen que esa relación no cambia en el "
              "tiempo, y ninguna se distinguió de las demás. Por eso la pregunta de este entregable es: **¿Un modelo que "
              "sigue día a día la tasa de suscripción puede predecir, antes de la llamada, qué clientes suscribirán un "
              "depósito a término mejor que los clasificadores estáticos?**")
    g.parrafo("La variable objetivo es y, que vale 1 si el cliente suscribe. Las predictoras son las 17 variables que se "
              "conocen antes de la llamada: el perfil del cliente, el canal de contacto, la historia de campañas y los "
              "indicadores macroeconómicos; duration queda por fuera porque solo se conoce al colgar. Es una clasificación "
              "binaria supervisada con **dependencia temporal**: los clientes llegan por días y las etiquetas de un día se "
              "conocen cuando ese día se cierra.")
    g.parrafo("Formalmente, el modelo original, que llamamos **L2E** (logística de estado latente en dos escalas), plantea "
              "P(yᵢ = 1 | θₜ) = σ(uₜ + xᵢᵀβₜ), donde uₜ es el nivel de conversión del día t, un estado latente que se mueve "
              "rápido, y βₜ son los efectos del cliente, que cambian despacio. El objetivo es maximizar el AUC-PR de "
              "predicciones hechas solo con información disponible, es decir, con las etiquetas de los días ya cerrados.")


def criterio_2(g, D):
    e, rango = D["eda"], D["rango"]
    sn = D["solo_nivel"].set_index("Puntaje")
    fuera = rango.loc["nr.employed", "Prueba fuera del rango de entrenamiento"]
    g.titulo("Guion", 2)
    g.parrafo(f"Los datos son Bank Marketing del repositorio UCI (bank-additional-full, licencia CC BY 4.0): {num(e['n_tr'] + e['n_te'], 0)} "
              "registros únicos, ordenados en el tiempo de mayo de 2008 a noviembre de 2010. La partición es cronológica: "
              f"el primer 80 % ({num(e['n_tr'], 0)} registros y {e['dias_tr']} días de contacto) para entrenar y el último "
              f"20 % ({num(e['n_te'], 0)} registros y {e['dias_te']} días) como prueba, usada una sola vez. La calidad es "
              "buena en forma, pero hay faltantes codificados como unknown, que son informativos, y un código especial "
              "(pdays = 999) que se trata como faltante con indicador.")
    g.parrafo("La prevención de fuga fue estricta. duration se excluye porque solo se conoce cuando la llamada termina: con "
              "ella el AUC de validación pasaría de 0,540 a 0,856. También hay **fuga temporal**: con validación aleatoria "
              "el AUC es de 0,670, y un modelo que solo conoce el año y el mes ya llega a 0,660; por eso toda la validación "
              "respeta el orden temporal. En el preprocesamiento se retiran euribor3m y emp.var.rate por multicolinealidad, "
              "se estandarizan las numéricas, se aplica log1p a campaign y previous y las categóricas van en one-hot, todo "
              "ajustado solo con el entrenamiento de cada pliegue.")
    g.parrafo("El hallazgo nuevo de este entregable es el que motiva el modelo: la tasa de suscripción de cada día es un "
              f"**nivel latente persistente**. Su logit tiene una autocorrelación de {num(e['acf1'], 2)} con el día anterior "
              f"y todavía de {num(e['acf5'], 2)} a cinco días, muy por encima de la banda de ruido (±{num(e['banda'], 2)}). Un "
              "puntaje que solo conoce la tasa de los días cerrados, sin ningún dato del cliente, alcanza en el bucle "
              f"externo el mismo AUC-PR que la mejor de las 140 combinaciones ({num(e['solo_nivel'])} frente a "
              f"{num(e['mejor_140'])}). Además, el {num(100 * fuera, 1)} % de la prueba tiene un nr.employed fuera del rango "
              "del entrenamiento, lo que castiga a los modelos que no extrapolan.")
    filas = []
    for nombre in sn.index:
        etiqueta = nombre.replace("Mejor benchmark: ", "Mejor de las 140 (") + (")" if nombre.startswith("Mejor") else "")
        filas.append([etiqueta] + [num(sn.loc[nombre, f"Pliegue {k}"]) for k in range(1, 6)] + [num(sn.loc[nombre, "Media"])])
    g.cuadro_("Puntaje que solo conoce la tasa de los días cerrados frente a la mejor combinación (AUC-PR por pliegue externo)",
              ["Puntaje", "P1", "P2", "P3", "P4", "P5", "Media"], filas,
              f"Fuente: cuaderno {CUAD}, sección 2.2 (Nuevo hallazgo: la tasa diaria de suscripción es un nivel latente "
              "persistente). Calculado solo con el entrenamiento, en los mismos cinco pliegues externos de los benchmarks.",
              anchos=[3960, 900, 900, 900, 900, 900, 900])


def criterio_3(g, D):
    p, c = D["params"], D["complejidad"]
    g.titulo("Guion", 2)
    g.parrafo("El L2E separa dos fuentes de variación. La **escala rápida** es el nivel del día: uₜ = uₜ₋₁ + εₜ, con "
              "εₜ ~ N(0, q_u·Δₜ), que recoge la lista de clientes, la campaña y el contexto de ese día. La **escala lenta** son "
              "las pendientes del cliente: βₜ = βₜ₋₁ + νₜ, con νₜ ~ N(0, q_β·Δₜ·I), que capturan la deriva de concepto. El "
              "modelo de observación es una logística: P(yᵢ = 1) = σ(uₜ + xᵢᵀβₜ).")
    g.parrafo("El aprendizaje tiene dos fases. En la **fase 1** se maximiza la log-posterior penalizada J(u, β) = Σ[yᵢηᵢ − "
              "log(1 + e^ηᵢ)] − Σ(uₖ − uₖ₋₁)²/(2·q_u·Δₖ) − (λ/2)‖β‖². Es la log-verosimilitud de Bernoulli más dos "
              "regularizaciones: la del paseo aleatorio, que suaviza el nivel, y la de tipo ridge, que encoge las pendientes. "
              "J es estrictamente cóncava, porque la log-verosimilitud es cóncava y las penalizaciones son definidas "
              "positivas, así que el máximo existe y es único. Se encuentra con Newton y búsqueda lineal de Armijo, que "
              "converge de forma global y cuadrática cerca del óptimo. El Hessiano tiene un bloque tridiagonal, el del nivel, "
              "que se elimina con el complemento de Schur; por eso cada iteración cuesta O(n·p² + T·p² + p³) y no "
              f"O((T + p)³). Con {num(c['T_max_denso'], 0)} días, la solución densa del mismo sistema tarda "
              f"{num(c['ratio_denso'], 1)} veces más.")
    g.parrafo("En la **fase 2**, en producción, el estado se propaga cada día, con P⁻ = P + Δ·diag(q_u, q_β, …), y cuando "
              "se cierran las etiquetas de un día se actualiza con la aproximación de Laplace: la moda y la curvatura de la "
              "posterior del día, que también es única porque el objetivo es estrictamente cóncavo. La probabilidad predictiva "
              "integra la incertidumbre del nivel con la aproximación probit, p̂ = σ(μ / √(1 + π·s²/8)). La predicción del "
              "día d solo usa etiquetas de días anteriores a d, lo que garantiza la ausencia de fuga por construcción.")
    g.parrafo("Los supuestos son cuatro: independencia condicional de los clientes dado el estado del día, dinámica de paseo "
              "aleatorio, pendientes constantes dentro de la ventana de entrenamiento y etiquetas disponibles con un día de "
              "demora. La capacidad de generalización viene de la separación de escalas: el nivel absorbe los cambios de "
              "prevalencia que un modelo estático no puede anticipar, y las pendientes se mantienen estables. El modelo, "
              "además, anida a los anteriores: con q_u y q_β iguales a cero es la logística L2, y la recalibración mensual "
              "que sugirió el profesor es su caso límite, con varianza a priori infinita, lote mensual y un mes de demora.")
    g.cuadro_("Hiperparámetros del L2E, su papel y su valor en el modelo final",
              ["Hiperparámetro", "Papel en el modelo", "Rango de búsqueda", "Valor final"],
              [["λ", "Precisión a priori de las pendientes (regularización ridge)", "0,1 a 1 000 (log)", num(p["lam"], 1)],
               ["q_u", "Varianza diaria del nivel (memoria de la escala rápida)", "0,001 a 1 (log)", num(p["q_nivel"], 3)],
               ["q_β", "Varianza diaria de las pendientes (escala lenta)", "10⁻⁷ a 10⁻² (log)",
                f"{p['q_pendiente']:.1e}".replace(".", ",")],
               ["Retardo L", "Días que tardan en conocerse las etiquetas", "Fijo por la operación", "1"]],
              f"Fuente: cuaderno {CUAD}, secciones 5.2 (formulación), 5.4 (propiedades) y 5.8 (optimización de "
              "hiperparámetros). El modelo final se optimizó con Optuna sobre todo el entrenamiento; q_β queda "
              f"{num(__import__('math').log10(p['q_nivel'] / p['q_pendiente']), 1)} órdenes de magnitud por debajo de q_u, "
              "lo que confirma con los datos la separación de escalas.",
              anchos=[1500, 4060, 2100, 1700])


def criterio_4(g, D):
    pu, hp, ab = D["pruebas_unitarias"], D["hiperparametros"], D["ablacion"].set_index("clave")
    hp5 = hp[hp["Pliegue"].isin([str(k) for k in range(1, 6)])]
    g.titulo("Guion", 2)
    g.parrafo("El modelo está implementado desde cero en la clase LogisticaDosEscalas, que hereda de BaseEstimator y "
              "ClassifierMixin de scikit-learn, con fit, predict, predict_proba y un método propio para la predicción "
              "secuencial. El cálculo numérico usa solo numpy y scipy, y no hay redes neuronales ni modelos importados. "
              f"Tiene {pu['total']} pruebas unitarias con pytest y pasan {pu['pasan']}: que aprende en datos simulados con "
              "solución conocida, reproducibilidad, entradas inválidas, compatibilidad con clone, ausencia de fuga y "
              "equivalencia con la logística L2 en el caso límite.")
    g.parrafo("Para que la comparación sea válida, el L2E pasa por **exactamente el mismo protocolo** que los 140 modelos: la "
              "misma partición cronológica, la misma validación anidada (cinco pliegues externos y tres internos con "
              "TimeSeriesSplit y una separación de un día), la misma semilla (42) y el mismo optimizador bayesiano, Optuna "
              "con TPE, con el mismo presupuesto de 20 configuraciones. En los pliegues de 2008 la validación interna eligió "
              f"un q_u pequeño, entre {num(hp5['q_u'].min(), 4)} y {num(hp5['q_u'].iloc[:4].max(), 4)}, y en el que ve el salto "
              f"de 2009 uno de {num(hp5['q_u'].iloc[4], 3)}: la optimización descubre la dinámica cuando los datos la contienen. "
              "El modelo final se reentrena con todo el entrenamiento y la prueba se evalúa una vez, con un día de demora "
              "en las etiquetas.")
    g.parrafo("Cada decisión de diseño se validó con una ablación, en la que cada componente se retira y la variante se "
              f"reoptimiza con el mismo protocolo. Sin la actualización en línea, el AUC-PR externo cae de "
              f"{num(ab.loc['l2e', 'AUC-PR externo'])} a {num(ab.loc['estatico', 'AUC-PR externo'])}, y con una sola escala, a "
              f"{num(ab.loc['escala_unica', 'AUC-PR externo'])}. En reproducibilidad, el modelo es determinista y la "
              f"semilla solo afecta al optimizador: con cuatro semillas, el AUC-PR de prueba va de {num(D['robustez']['sem_min'])} "
              f"a {num(D['robustez']['sem_max'])}. Todo se ejecuta con los scripts del repositorio y las versiones fijadas en "
              "requirements.txt.")
    filas = [[ab.loc[k, "Variante"], num(ab.loc[k, "AUC-PR externo"]), num(ab.loc[k, "AUC-PR prueba"]),
              num(ab.loc[k, "Brier prueba"])] for k in ab.index]
    g.cuadro_("Estudio de ablación del L2E (cada variante reoptimizada con la misma validación anidada)",
              ["Variante", "AUC-PR externo", "AUC-PR prueba", "Brier prueba"], filas,
              f"Fuente: cuaderno {CUAD}, sección 5.9 (Estudio de ablación). C1: nivel latente en el ajuste; C2: actualización "
              "del nivel; C3: actualización de las pendientes; C4: integración de la incertidumbre.",
              anchos=[4560, 1600, 1600, 1600])


def criterio_5(g, D):
    r3, res = D["r3"], D["resumen"].set_index("clave")
    pt = D["pruebas_r3"]
    rob = D["robustez"]
    mejor = r3["mejor"]
    rk = D["adicionales_rangos"].set_index("Modelo")
    adic = [c for c in rk.columns if c != "bank marketing"]
    c = D["correcciones"]
    g.titulo("Guion", 2)
    g.parrafo("Antes de los resultados, la retroalimentación del primer entregable. El docente descontó por el cambio de "
              "prevalencia, la inestabilidad de los coeficientes con C = 100 y el componente temporal parcial, y pidió "
              "recalibrar el intercepto y comparar modelos más flexibles. Todo quedó resuelto en el Entregable 2 y se "
              f"mantiene en este (cuadro {g.cuadro + 1}): la probabilidad media del L2E en la prueba es de "
              f"{num(c['prob_l2e'])}, igual a la tasa observada; los coeficientes varían entre pliegues "
              f"{num(c['de_c100'] / c['de_l2e'], 1)} veces menos que con C = 100, y el umbral elegido con el pasado da un "
              f"F1 de {num(c['f1_val'])} en la prueba, frente a 0,383 en el primer entregable.")
    filas = [
        ["Cambio de prevalencia (probabilidades que la subestiman)", "Probabilidad media de 0,243 frente a una tasa de 0,308",
         "Recalibración mensual del intercepto: probabilidad media de 0,300",
         f"Nivel diario del L2E: probabilidad media de {num(c['prob_l2e'])}; error mensual de {num(c['err_l2e'])} frente "
         f"a {num(c['err_sin'])} sin recalibrar"],
        ["Inestabilidad de los coeficientes con C = 100", "C = 100 elegido por el AUC de validación",
         "C = 0,01 por pérdida logarítmica: desviación bootstrap de 0,250 a 0,053",
         f"Desviación mediana entre pliegues: {num(c['de_c100'])} (C = 100), {num(c['de_c001'])} (C = 0,01) y "
         f"{num(c['de_l2e'])} (L2E, λ = {num(c['lam'], 1)})"],
        ["Componente temporal parcial", "Solo contexto económico y validación aleatoria frente a temporal",
         "Serie diaria, régimen y deriva de concepto: AUC de 0,667 aleatoria frente a 0,514 temporal",
         f"El nivel diario entra al modelo (autocorrelación de {num(D['eda']['acf1'], 2)}); validación anidada temporal "
         "y evaluación precuencial"],
        ["Recalibrar el intercepto con datos recientes", "Sin recalibración", "Con el mes anterior cerrado: AUC de 0,710 a 0,752",
         "El nivel del L2E es un intercepto que se recalibra cada día; la recalibración mensual es un caso límite (P5)"],
        ["Comparar modelos más flexibles", "Solo la logística frente al Dummy",
         "Splines, bosque y boosting: AUC de 0,751, 0,735 y 0,697 frente a 0,752",
         "140 combinaciones de 7 modelos; los árboles pierden porque el "
         f"{num(100 * D['rango'].loc['nr.employed', 'Prueba fuera del rango de entrenamiento'], 1)} % de la prueba sale "
         "del rango de nr.employed"],
        ["AUC moderado y umbral que empeora en prueba", "AUC de 0,677; el umbral de 0,60 baja el F1 de 0,428 a 0,383",
         "Umbral de 0,11: F1 de 0,157 a 0,525; recalibrado (0,13), 0,563",
         f"AUC de {num(res.loc['L2E', 'AUC prueba'])}; umbral de {num(c['u_val'], 2)} elegido sin la prueba: F1 de "
         f"{num(c['f1_val'])} (óptimo a posteriori: {num(c['f1_opt'])})"]]
    g.cuadro_("Retroalimentación del primer entregable: cómo quedó cada observación",
              ["Observación del docente", "Entregable 1 (calificado)", "Corrección (Entregable 2)", "Cuarto entregable"],
              filas,
              f"Fuente: cuaderno {CUAD}, sección 1.2 (tabla de trazabilidad de correcciones), con la evidencia de los "
              "cuadernos «Entregable 1 - Bank Marketing.ipynb» (secciones 4.1, 4.3, 4.4 y 4.7) y «Entregable 2 - Bank "
              "Marketing.ipynb» (secciones 3.6.3 a 3.6.5, 4.3, 4.7.1, 4.9.1 y 4.11).",
              anchos=[2000, 2200, 2460, 2700])
    g.parrafo("La métrica principal es el AUC-PR, porque la clase positiva es minoritaria, el uso real es ordenar a quién "
              "llamar y no depende del umbral; se reportan además el AUC y el Brier, que mide la calidad de las "
              f"probabilidades. En el bucle externo, el L2E obtiene {num(r3['ext_l2e'])} frente a {num(r3['ext_mejor'])} de la "
              f"mejor combinación del tercer entregable. En la prueba llega a {num(res.loc['L2E', 'AUC-PR prueba'])} de AUC-PR y "
              f"{num(res.loc['L2E', 'AUC prueba'])} de AUC, frente a {num(res.loc['referencia_e2', 'AUC-PR prueba'])} y "
              f"{num(res.loc['referencia_e2', 'AUC prueba'])} de la Referencia E2, y su Brier es el más bajo "
              f"({num(res.loc['L2E', 'Brier prueba'])}).")
    g.parrafo(f"La superioridad se verificó con pruebas de hipótesis. Frente a la mejor combinación, {r3['n_r3']} de las 4 "
              "pruebas declaradas rechazan la igualdad con p ajustado por Holm menor que 0,05 (Wilcoxon por meses, bootstrap "
              f"y permutación por días), con un δ de Cliff de {num(r3['cliff'])}. Por eso se cumple la condición R3 de la guía. "
              "Como DeLong supone independencia y los clientes de un mismo día comparten el nivel, las pruebas válidas aquí "
              "son las que remuestrean días completos.")
    g.parrafo("El análisis de errores es honesto. Dentro de cada mes, el L2E no ordena mejor que las logísticas ni que la "
              "SVM: su ventaja viene de seguir el nivel entre días, no de distinguir mejor a los clientes de un mismo día. "
              f"Supera a {NOMBRES[mejor]} en "
              f"{rob['meses_gana_mejor']} de {rob['meses']} meses, y en los tres conjuntos adicionales no es el mejor "
              f"modelo (puestos {', '.join(str(int(rk.loc['L2E · completo', c])) for c in adic)}). No hay sobreajuste "
              f"a una configuración: el {num(100 * rob['frac_q'], 0)} % de la rejilla de hiperparámetros supera al mejor "
              f"benchmark, y aun con 40 días de demora en las etiquetas el AUC-PR es de {num(rob['retardo_40'])}.")
    filas = []
    for k in res.index:
        filas.append([res.loc[k, "Modelo"].replace(" · ", ", ", 2), res.loc[k, "AUC-PR externo"],
                      f"{num(res.loc[k, 'AUC-PR prueba'])} {res.loc[k, 'IC 95 % AUC-PR']}", num(res.loc[k, "AUC prueba"]),
                      "—" if pd.isna(res.loc[k, "Brier prueba"]) else num(res.loc[k, "Brier prueba"])])
    g.cuadro_("Modelo original frente a los benchmarks: bucle externo y prueba",
              ["Modelo", "AUC-PR externo", "AUC-PR prueba [IC 95 %]", "AUC prueba", "Brier prueba"], filas,
              f"Fuente: cuaderno {CUAD}, sección 6.2 (Tabla resumen). Intervalos por bootstrap estacionario de días con 1 000 "
              "réplicas. Los benchmarks son la mejor combinación de cada modelo en el tercer entregable; la SVM no tiene "
              "probabilidades y la Referencia E2 no tiene bucle externo.", anchos=[3460, 1500, 2000, 1100, 1300])
    filas = []
    for _, r in pt.iterrows():
        filas.append([r["Prueba"], r["Estadístico"], pv(r["p ajustado"]) if pd.notna(r["p ajustado"]) else "—",
                      r["Efecto"].replace("d̄", "Δ medio"), r["Conclusión"]])
    g.cuadro_(f"Pruebas de hipótesis del L2E frente a la mejor combinación benchmark ({NOMBRES[mejor]})",
              ["Prueba", "Estadístico", "p ajustado", "Efecto", "Conclusión"], filas,
              f"Fuente: cuaderno {CUAD}, sección 6.7 (Tabla de pruebas de hipótesis). p ajustado por Holm sobre las 8 "
              "comparaciones del L2E; las cuatro primeras pruebas son las declaradas para la condición R3.",
              anchos=[2860, 1500, 1100, 2200, 1700])


def criterio_6(g, D):
    r3 = D["r3"]
    g.titulo("Guion", 2)
    g.parrafo("En síntesis, cuatro conclusiones. Primera: el límite de los 140 modelos del tercer entregable era temporal, "
              "no de flexibilidad; la tasa diaria es un nivel latente persistente que ningún modelo estático representa. "
              "Segunda: el L2E convierte ese hallazgo en un modelo, con dos escalas de tiempo, un ajuste por bloques en "
              "tiempo lineal y un filtro de Laplace con etiquetas diferidas, y la recalibración mensual que sugirió el "
              f"profesor es su caso límite. Tercera: cumple R1, R2 y R3, con {r3['n_r3']} pruebas significativas frente a la "
              "mejor combinación y probabilidades calibradas. Cuarta: el mecanismo es transferible; el filtro de nivel mejora "
              "a la mayoría de los modelos de referencia cuando se les aplica.")
    g.parrafo("Las limitaciones son igual de importantes. La significancia depende de la unidad de análisis: en 2008 ningún "
              "modelo discrimina, así que la evidencia viene de la prueba. Dentro de cada mes, el L2E no supera a una buena "
              "logística. El modelo necesita que las etiquetas lleguen pronto, y en los tres conjuntos adicionales no es el "
              "mejor, porque allí el nivel latente es débil o lo explican las variables. Además, se declara que durante el "
              "prototipado se consultó la prueba con variantes exploratorias.")
    g.parrafo("Por ello, el trabajo futuro es validar con datos posteriores a 2010 el ensamble con un bosque reentrenado cada "
              "mes, que mejoró la prueba pero no la validación anidada; buscar una forma de elegir q_u con ventanas de ajuste "
              "cortas, e integrar el L2E al dashboard y al Jupyter Book.")


PREGUNTAS = [
    ("¿Cómo atendió la retroalimentación del primer entregable?",
     "Las cinco observaciones quedaron resueltas en el Entregable 2: C = 0,01 elegido por pérdida logarítmica, recalibración "
     "mensual del intercepto, análisis temporal completo y comparación con splines, bosque y boosting. En este entregable se "
     "llevan al modelo: el nivel del L2E es un intercepto que se recalibra cada día, su probabilidad media en la prueba es "
     "{prob} frente a una tasa de {prev}, y el umbral elegido con el pasado ya no empeora en la prueba (F1 de {f1})."),
    ("¿Por qué un modelo propio y no otro clasificador más potente?",
     "Porque el diagnóstico mostró que el problema era temporal: con validación aleatoria los árboles sí ganan, pero en la "
     "prueba pierden porque el {fuera} % de los registros tiene nr.employed fuera del rango de entrenamiento. Un clasificador "
     "más flexible no resuelve eso; un nivel latente que se actualiza, sí."),
    ("¿No es trampa usar las etiquetas de los días anteriores?",
     "No, porque el modelo solo usa días ya cerrados, como haría el banco en producción, y lo garantiza por construcción: una "
     "prueba unitaria cambia las etiquetas desde un día y verifica que las predicciones anteriores no cambian. Además, se "
     "comparó contra los benchmarks con la misma información y el L2E sigue adelante."),
    ("¿En qué se diferencia de un filtro de Kalman o de una logística dinámica?",
     "El Kalman supone observaciones gaussianas; aquí son de Bernoulli y se actualizan con Laplace. La logística dinámica "
     "clásica usa una sola tasa de olvido; el L2E separa el nivel rápido de las pendientes lentas, y con una sola escala el "
     "AUC-PR externo cae a {escala_unica}."),
    ("¿Qué pasa si las etiquetas llegan tarde?",
     "La ventaja se reduce pero no desaparece: con un día de demora el AUC-PR de prueba es {ret1} y con 40 días, {ret40}, "
     "todavía por encima del mejor benchmark estático."),
    ("Si Friedman no rechaza, ¿por qué dicen que es superior?",
     "Porque Friedman compara rangos dentro de cada mes, y la ventaja del L2E está en ordenar entre días. La guía define R3 "
     "frente a la mejor combinación: {n_r3} de las 4 pruebas declaradas rechazan la igualdad y el efecto es pequeño "
     "(δ de Cliff de {cliff})."),
    ("¿Es mejor que la logística del Entregable 2?",
     "En el ordenamiento global de la prueba sí, con bootstrap y permutación por días significativos. Dentro de cada mes no: "
     "la logística ordena igual de bien a los clientes de un mismo periodo. Su aporte es el nivel, no las pendientes."),
    ("Si la variante sin actualizar las pendientes tiene mejor AUC-PR externo, ¿por qué no la eligieron?",
     "Porque la diferencia no se distingue del ruido: es de {sp_dif}, gana {sp_gana} de los 5 pliegues, la t corregida de "
     "Nadeau-Bengio da p = {sp_p} y la prueba bayesiana asigna {sp_rope} de probabilidad a la equivalencia práctica. La "
     "ablación mide el aporte de cada componente, no elige el modelo; con el bucle externo empatado se conserva el modelo "
     "declarado antes de experimentar, y la prueba no interviene en esa decisión."),
    ("¿Por qué no adoptaron el ensamble M4 si mejora la prueba?",
     "Porque la regla declarada decide con el bucle externo, y ahí M4 empeora. Elegirlo por su resultado en la prueba sería "
     "usar la prueba para seleccionar el modelo, que es justo lo que la guía sanciona; queda como línea futura."),
    ("¿Cómo saben que el modelo es nuevo?",
     "Por la revisión sistemática: {bases} bases, {consultas} consultas y {revisados} registros revisados. Cada componente "
     "tiene literatura, pero la intersección con el telemercadeo bancario está vacía; el único antecedente cercano usa "
     "árboles con actualización incremental, otra familia."),
    ("¿Por qué converge el ajuste?",
     "Porque el objetivo es estrictamente cóncavo: la log-verosimilitud de Bernoulli es cóncava y las penalizaciones son "
     "definidas positivas. Newton con búsqueda lineal converge al único máximo, y la prueba unitaria verifica que el objetivo "
     "nunca decrece."),
    ("¿Por qué el AUC-PR y no el accuracy?",
     "Porque la clase positiva es minoritaria y el objetivo es ordenar a quién llamar. Decir «no» a todos da un 69 % de "
     "accuracy en la prueba; el AUC-PR, en cambio, tiene como línea base la prevalencia, {prev}."),
    ("¿Qué harían con más tiempo?",
     "Validar el ensamble M4 con datos posteriores a 2010, buscar una forma de elegir q_u con ventanas cortas, porque el "
     "Bayes empírico falló ante el salto de 2009, y extender el nivel a grupos, por canal o por campaña."),
]


def criterio_7(g, D):
    r3, rob, lit = D["r3"], D["robustez"], D["literatura"]
    ab = D["ablacion"].set_index("clave")
    sp = D["ablacion_sp"]
    v = {"prob": num(D["correcciones"]["prob_l2e"]), "f1": num(D["correcciones"]["f1_val"]),
         "sp_dif": num(sp["dif"], 4), "sp_gana": sp["gana"], "sp_p": num(sp["p_nb"], 2), "sp_rope": num(sp["p_rope"], 2),
         "fuera": num(100 * D["rango"].loc["nr.employed", "Prueba fuera del rango de entrenamiento"], 1),
         "escala_unica": num(ab.loc["escala_unica", "AUC-PR externo"]), "ret1": num(rob["retardo_1"]),
         "ret40": num(rob["retardo_40"]), "n_r3": r3["n_r3"], "cliff": num(r3["cliff"]), "bases": lit["bases"],
         "consultas": lit["consultas"], "revisados": num(lit["revisados"], 0), "prev": num(D["eda"]["prev_te"])}
    g.parrafo("Este criterio no se expone: se evalúa con las preguntas del jurado. Como el proyecto es individual, el "
              "dominio se demuestra respondiendo con precisión y con las cifras del propio trabajo. Cada respuesta debe "
              "tomar entre 20 y 40 segundos: primero la respuesta directa, después la cifra que la respalda y al final la "
              "decisión que se tomó.")
    g.titulo("Preguntas probables y respuestas preparadas", 2)
    for pregunta, respuesta in PREGUNTAS:
        g.parrafo(f"**{pregunta}**")
        g.parrafo(respuesta.format(**v))


APOYO = {
    1: "Cuaderno del cuarto entregable, secciones 0 (resumen ejecutivo) y 5.2 (formulación).",
    2: "Cuaderno, secciones 2.1 y 2.2 (partición y nivel diario latente) y figura de la tasa diaria.",
    3: "Cuaderno, secciones 5.2 a 5.4 (formulación, pseudocódigo y propiedades). Conviene tener a mano las ecuaciones.",
    4: "Cuaderno, secciones 5.5 a 5.9 (código, pruebas unitarias, complejidad, optimización y ablación).",
    5: "Cuaderno, secciones 1.2 (trazabilidad de correcciones) y 6.2 a 6.7 (tabla resumen, diagramas de diferencia "
       "crítica y tabla de pruebas).",
    6: "Cuaderno, secciones 6.12 y 8 (veredicto, conclusiones y limitaciones).",
    7: "Ninguno: se responde de forma oral. Conviene tener abiertos el cuaderno y el guion por si piden una cifra.",
}


def palabras(doc, k):
    """Palabras del guion del criterio k (párrafos normales entre su título y el siguiente criterio)."""
    body = list(doc.element.body.iterchildren())
    i = next(j for j, el in enumerate(body) if _es_h1(doc, el) and _texto(el).startswith(f"{k}. "))
    n = 0
    for el in body[i + 2:]:
        if el.tag == qn("w:sectPr") or _es_h1(doc, el):
            break
        if el.tag == qn("w:p"):
            ps = el.find(f".//{qn('w:pStyle')}")
            if ps is None:
                n += len(re.findall(r"\w+", _texto(el)))
    return n


def main():
    D = cargar()
    doc = docx.Document(GUION)
    # Encabezado e introducción
    pars = doc.paragraphs
    _poner_texto(pars[1]._p, "Proyecto final de Machine Learning · Cuarto entregable: modelo original L2E · Bank Marketing (UCI)")
    _poner_texto(pars[3]._p, "Este documento resuelve, en el orden de la rúbrica, los siete criterios de evaluación de la "
                 "sustentación del cuarto entregable, cuyo componente central es el modelo original L2E. Para cada uno se "
                 "presentan los aspectos que evalúa el jurado, el guion para exponerlo y el material que lo apoya. La "
                 "extensión de cada guion está calibrada con el tiempo que la rúbrica asigna al criterio, a un ritmo de unas "
                 "130 palabras por minuto, de forma que la exposición completa ocupe los 20 minutos previstos. Todas las "
                 "cifras provienen del cuaderno ejecutado «Proyecto Final - Cuarto Entregable - Bank Marketing.ipynb». La "
                 "sección 8 documenta las mejoras y los ajustes evaluados después del modelo original.")
    i_rec = next(i for i, p in enumerate(pars) if p.text.startswith("Recomendación de uso"))
    _poner_texto(pars[i_rec]._p, "Recomendación de uso: el criterio 3 pesa el 25 % de la nota y tiene 5 minutos; conviene "
                 "ensayarlo con cronómetro, porque las ecuaciones del L2E se prestan a excederse. En los demás criterios, si "
                 "el tiempo apremia, se puede omitir el último párrafo de cada guion sin perder el hilo.")
    escritores = {1: criterio_1, 2: criterio_2, 3: criterio_3, 4: criterio_4, 5: criterio_5, 6: criterio_6, 7: criterio_7}
    # Se limpian todos los criterios antes de escribir, para que la numeración de cuadros empiece en 1
    recuadros = {k: limpiar_criterio(doc, k) for k in escritores}
    # La sección 8 se elimina aquí y la regenera actualizar_guion.py con la numeración consecutiva
    from guion_docx import quitar_seccion
    quitar_seccion(doc, "8. Mejoras y ajustes del modelo")
    g = Guion(doc)
    for k, escribir in escritores.items():
        _celda(recuadros[k], 2, 1, APOYO[k])
        g.desde(recuadros[k])
        escribir(g, D)
    g.ancla = None
    g.guardar(GUION)
    doc = docx.Document(GUION)
    tabla1 = doc.tables[0]._tbl
    for k in range(1, 7):
        _celda(tabla1, k, 3, f"≈ {palabras(doc, k)}")
    doc.save(GUION)
    print("Criterios 1 a 7 reescritos;", {k: palabras(doc, k) for k in range(1, 7)}, "palabras por criterio")


if __name__ == "__main__":
    main()
